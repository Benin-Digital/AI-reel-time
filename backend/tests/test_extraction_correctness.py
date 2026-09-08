"""Tests de non-regression pour des bugs d'extraction trouves par inspection
+ reproduction manuelle (pas de test existant ne couvrait ces chemins) :

- DOCX : l'ordre paragraphes/tableaux n'etait pas preserve (doc.paragraphs et
  doc.tables sont deux listes separees en python-docx qui ne refletent pas
  l'ordre reel du document -> tous les tableaux finissaient a la fin).
- DOCX : le texte des en-tetes/pieds de page (souvent le nom/contact du
  candidat dans un CV) etait entierement ignore.
- TXT : le fallback latin-1 ne leve jamais d'erreur et masque les fichiers
  encodes en cp1252 (guillemets typographiques/tirets Windows tres courants),
  produisant des caracteres de controle a la place.
- clean_text (partage PDF/DOCX/TXT) : la deduplication plafonnait a 2
  occurrences n'importe quelle ligne, supprimant du contenu legitime repete
  (ex: une meme competence citee dans 3+ experiences differentes).
"""
from __future__ import annotations

import docx

from app.services.extraction import clean_text, extract_text_from_docx, extract_text_from_txt


def test_docx_preserves_paragraph_table_order(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Titre: Développeur Python")
    doc.add_paragraph("Expérience:")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "ACME Corp"
    table.cell(0, 1).text = "2019-2022"
    table.cell(1, 0).text = "Backend dev"
    table.cell(1, 1).text = "Python/Django"
    doc.add_paragraph("Formation: Master info 2018")

    path = tmp_path / "cv.docx"
    doc.save(path)

    result = extract_text_from_docx(path)
    lines = result.splitlines()

    assert lines.index("Titre: Développeur Python") < lines.index("Expérience:")
    table_line_idx = next(i for i, l in enumerate(lines) if "ACME Corp" in l)
    assert lines.index("Expérience:") < table_line_idx, "le tableau doit rester apres le texte qui le precede"
    assert table_line_idx < lines.index("Formation: Master info 2018"), \
        "le tableau ne doit pas etre repousse apres du texte qui le suit reellement dans le document"


def test_docx_captures_header_and_footer_text(tmp_path):
    doc = docx.Document()
    doc.sections[0].header.paragraphs[0].text = "Jean Dupont - jean.dupont@email.com - 06 12 34 56 78"
    doc.sections[0].footer.paragraphs[0].text = "CV genere le 01/01/2026"
    doc.add_paragraph("Corps du document normal.")

    path = tmp_path / "cv_with_header.docx"
    doc.save(path)

    result = extract_text_from_docx(path)

    assert "Jean Dupont" in result
    assert "jean.dupont@email.com" in result
    assert "Corps du document normal." in result
    assert "CV genere le 01/01/2026" in result


def test_txt_decodes_windows_cp1252_smart_quotes(tmp_path):
    text = "Poste : Développeur 'Senior' – Paris"
    path = tmp_path / "job.txt"
    path.write_bytes(text.encode("cp1252"))

    result = extract_text_from_txt(path)

    assert "Senior" in result
    assert "\x91" not in result and "\x92" not in result and "\x96" not in result, \
        "le texte ne doit pas contenir de caracteres de controle (mojibake latin-1)"


def test_clean_text_keeps_short_line_repeated_more_than_twice():
    text = (
        "Experience 1 - ACME Corp\n- Python\n- Django\n\n"
        "Experience 2 - Beta Inc\n- Python\n- Flask\n\n"
        "Experience 3 - Gamma SAS\n- Python\n- FastAPI\n\n"
        "Experience 4 - Delta Ltd\n- Python\n- Docker"
    )

    result = clean_text(text)

    assert result.count("Python") == 4, "une competence citee dans 4 experiences differentes ne doit pas etre tronquee a 2"


def test_clean_text_still_dedupes_long_repeated_boilerplate():
    header = "Curriculum Vitae confidentiel - Jean Dupont - ne pas diffuser"
    text = "\n".join([header, "Page 1", header, "Page 2", header, "Page 3", header])

    result = clean_text(text)

    assert result.count(header) == 2, "un en-tete/pied de page long repete sur chaque page doit toujours etre plafonne"
